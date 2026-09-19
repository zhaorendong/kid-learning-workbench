# -*- coding: utf-8 -*-
"""校验生成的 Word 题库：结构、网格数量、边框设置、内容抽样"""
import os
from collections import Counter

from docx import Document
from docx.oxml.ns import qn

PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                    '一年级数独题库（可打印）.docx')


def walk(tables, depth=0):
    for t in tables:
        yield depth, t
        for row in t.rows:
            for cell in row.cells:
                yield from walk(cell.tables, depth + 1)


def cell_border(cell, tag):
    tcPr = cell._tc.find(qn('w:tcPr'))
    if tcPr is None:
        return None
    bs = tcPr.find(qn('w:tcBorders'))
    if bs is None:
        return None
    e = bs.find(qn('w:' + tag))
    if e is None:
        return None
    return (e.get(qn('w:val')), int(e.get(qn('w:sz'))))


def main():
    assert os.path.exists(PATH), '文件不存在: ' + PATH
    doc = Document(PATH)
    print('文件大小: %.1f KB' % (os.path.getsize(PATH) / 1024))

    top = doc.tables
    print('顶层表格数:', len(top))

    grid_counter = Counter()
    layout_depth1 = 0
    bad = []

    for depth, t in walk(top):
        r, c = len(t.rows), len(t.columns)
        if depth == 0:
            continue
        if r == c and r in (4, 6, 9) and depth == 1:
            grid_counter[r] += 1
            # 抽查：左上角外框应为粗线（sz>=12），格间应为细线（sz<=8）
            tl = t.rows[0].cells[0]
            for tag in ('top', 'left'):
                b = cell_border(tl, tag)
                if not b or b[1] < 12:
                    bad.append(('左上外框 ' + tag, b))
            if t.rows[0].cells[1]._tc is not None:
                b = cell_border(t.rows[0].cells[1], 'left')
                if b and b[1] >= 12:
                    bad.append(('格间线应为细线 ' + str(b), ))
        elif depth == 1:
            layout_depth1 += 1

    print('数独网格统计:', dict(grid_counter))
    print('边框异常:', bad[:5] if bad else '无')

    total_grids = sum(grid_counter.values())
    print('网格总数:', total_grids, '(期望 90 题 + 90 答案 = 180)')

    # 抽样内容
    shown = {'puzzle': 0, 'answer': 0}
    for depth, t in walk(top):
        if depth == 1 and len(t.rows) == 4:
            kind = 'puzzle' if shown['puzzle'] == 0 else 'answer'
            if shown[kind] >= 2:
                break
            shown[kind] += 1
            print()
            print('--- 4x4 网格抽样 (%s) ---' % kind)
            for row in t.rows:
                vals = []
                for cell in row.cells:
                    txt = ''.join(p.text for p in cell.paragraphs).strip()
                    vals.append(txt if txt else '.')
                print('   ', ' '.join(vals))

    # 题号抽样（在布局表格的 cell 段落里）
    titles = []
    for depth, t in walk(top):
        if depth == 0:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.text.strip().startswith('第'):
                            titles.append(p.text.strip())
    print()
    print('题号标签抽样:', titles[:3], '...', titles[-3:] if titles else '')
    print('题号标签总数:', len(titles), '(期望 90)')
    return 0 if not bad else 1


if __name__ == '__main__':
    raise SystemExit(main())
